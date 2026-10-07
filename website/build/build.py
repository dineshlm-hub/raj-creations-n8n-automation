"""Builds the rajcreation.info static site into website/public/.

    python build.py            # full build (re-renders catalogue images only if missing)
    python build.py --images   # force re-rendering every catalogue image

Inputs:
    ../../*.pdf                               the three catalogue PDFs (see catalogue_pages.py)
    ../Website Rates (for Google Sheet).xlsx  fallback prices, used until the Google Sheet is connected
    config.py                                 secret price-page token, Google Sheet CSV link, phone numbers
"""
import hashlib
import html
import io
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pymupdf
from PIL import Image

import config
from catalogue_pages import CATEGORIES, PAGES

HERE = Path(__file__).resolve().parent
SITE = HERE.parent
ROOT = SITE.parent
SRC = SITE / "src"
OUT = SITE / "public"
RATES_XLSX = SITE / "Website Rates (for Google Sheet).xlsx"

THUMB_W, FULL_W = 480, 1200


def norm(s):
    return re.sub(r"[^A-Z0-9]", "", str(s).upper())


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")


def render_images(force):
    """PDF page -> img/<cat>/<page>-t.webp (thumbnail) and -f.webp (full)."""
    sizes = {}
    for cat in CATEGORIES:
        doc = pymupdf.open(ROOT / cat["pdf"])
        out_dir = OUT / "img" / cat["id"]
        out_dir.mkdir(parents=True, exist_ok=True)
        for page_no, _, _ in PAGES[cat["id"]]:
            full, thumb = out_dir / f"{page_no}-f.webp", out_dir / f"{page_no}-t.webp"
            page = doc[page_no - 1]
            w, h = page.rect.width, page.rect.height
            sizes[(cat["id"], page_no)] = (FULL_W, round(FULL_W * h / w))
            if full.exists() and thumb.exists() and not force:
                continue
            pix = page.get_pixmap(matrix=pymupdf.Matrix(FULL_W / w, FULL_W / w))
            im = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
            im.save(full, "WEBP", quality=72, method=6)
            t = im.resize((THUMB_W, round(THUMB_W * im.height / im.width)), Image.LANCZOS)
            t.save(thumb, "WEBP", quality=68, method=6)
            print(f"  rendered {cat['id']} p{page_no}")
    return sizes


def render_logo():
    """Crop the Raj Creation logo from the top-left of the first clutchers page."""
    out = OUT / "assets" / "logo.webp"
    if out.exists():
        return
    doc = pymupdf.open(ROOT / CATEGORIES[0]["pdf"])
    page = doc[0]
    clip = pymupdf.Rect(0, 0, page.rect.width * 0.26, page.rect.height * 0.105)
    pix = page.get_pixmap(clip=clip, matrix=pymupdf.Matrix(4, 4))
    im = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
    # trim white border
    bg = Image.new("RGB", im.size, (255, 255, 255))
    from PIL import ImageChops
    box = ImageChops.difference(im, bg).convert("L").point(lambda p: 255 if p > 24 else 0).getbbox()
    if box:
        pad = 8
        im = im.crop((max(box[0] - pad, 0), max(box[1] - pad, 0), min(box[2] + pad, im.width), min(box[3] + pad, im.height)))
    im.thumbnail((360, 200))
    out.parent.mkdir(parents=True, exist_ok=True)
    im.save(out, "WEBP", quality=90)


def render_logo_mark():
    """The RC monogram from the catalogue logo, on a transparent background (the name is set as text on the site)."""
    out = OUT / "assets" / "logo-mark.webp"
    if out.exists():
        return
    img = monogram()
    img.thumbnail((104, 90), Image.LANCZOS)  # shown about 52 x 42 px: twice that is sharp on phones
    img.save(out, "WEBP", quality=82, method=6)


def render_logo_pdf():
    """Logo for the bot's PDF order confirmation: the RC monogram with the name set underneath
    (the catalogue logo spells the name "Raj Creations"; the company name is "Raj Creation")."""
    from PIL import ImageDraw, ImageFont
    out = OUT / "assets" / "logo-pdf.jpg"
    mark = monogram()
    mark.thumbnail((400, 230), Image.LANCZOS)
    font = ImageFont.truetype(str(SRC / "assets" / "fonts" / "playfair-latin.woff2"), 62)
    font.set_variation_by_name("SemiBold")
    name, spacing, ink = "RAJ CREATION", 6, (0x24, 0x1c, 0x22)
    widths = [font.getlength(ch) for ch in name]
    text_w = sum(widths) + spacing * (len(name) - 1)
    pad, gap, text_h = 24, 22, 62
    w = round(max(mark.width, text_w) + 2 * pad)
    h = pad + mark.height + gap + text_h + 18 + pad
    im = Image.new("RGB", (w, h), (255, 255, 255))
    im.paste(mark, ((w - mark.width) // 2, pad), mark)
    d = ImageDraw.Draw(im)
    x, y = (w - text_w) / 2, pad + mark.height + gap
    for ch, cw in zip(name, widths):
        d.text((x, y), ch, font=font, fill=ink)
        x += cw + spacing
    line_y = y + text_h + 10
    d.line(((w - text_w) / 2, line_y, (w + text_w) / 2, line_y), fill=ink, width=3)
    im.save(out, "JPEG", quality=90)
    return out


def monogram():
    """The RC monogram cut out of the catalogue logo at full resolution (RGBA, ink on transparent)."""
    import numpy as np
    from scipy import ndimage
    doc = pymupdf.open(ROOT / CATEGORIES[0]["pdf"])
    page = doc[0]
    clip = pymupdf.Rect(0, 0, page.rect.width * 0.26, page.rect.height * 0.105)
    pix = page.get_pixmap(clip=clip, matrix=pymupdf.Matrix(8, 8))
    a = np.asarray(Image.open(io.BytesIO(pix.tobytes("png"))).convert("L")).astype(float)
    bg = np.percentile(a, 90)
    ink = np.clip((bg - a) / (bg - 40), 0, 1)
    ink[(ink > 0.5).mean(axis=1) > 0.35] = 0                     # the underline under the name
    ys, xs = np.where(ink > 0.15)
    ink = ink[ys.min():ys.max() + 1, xs.min():xs.max() + 1]
    prof = (ink > 0.3).sum(axis=1)
    h = ink.shape[0]
    split = min(range(int(h * 0.6), int(h * 0.8)), key=lambda y: prof[y])  # gap above "RAJ CREATIONS"
    mark = ink[:split].copy()
    mark[mark < 0.22] = 0
    lab, n = ndimage.label(mark > 0.3)
    keep = np.zeros(n + 1, bool)
    for i in range(1, n + 1):
        yy, _ = np.where(lab == i)
        keep[i] = not (yy.mean() > mark.shape[0] * 0.88 and len(yy) < 2000)  # bits of the name's first letter
    mask = keep[lab] | (ndimage.binary_dilation(keep[lab], iterations=2) & (mark > 0))
    mark = np.clip((np.where(mask, mark, 0) - 0.22) / 0.6, 0, 1)
    ys, xs = np.where(mark > 0.05)
    mark = mark[ys.min():ys.max() + 1, xs.min():xs.max() + 1]
    rgba = np.zeros(mark.shape + (4,), dtype=np.uint8)
    rgba[..., 0], rgba[..., 1], rgba[..., 2] = 0x24, 0x1c, 0x22
    rgba[..., 3] = (mark * 255).astype(np.uint8)
    return Image.fromarray(rgba, "RGBA")


def build_catalogue(sizes):
    cats = []
    for cat in CATEGORIES:
        items = []
        for page_no, codes, note in PAGES[cat["id"]]:
            img = {
                "t": f"img/{cat['id']}/{page_no}-t.webp",
                "f": f"img/{cat['id']}/{page_no}-f.webp",
                "w": sizes[(cat["id"], page_no)][0],
                "h": sizes[(cat["id"], page_no)][1],
                "page": page_no,
            }
            if items and items[-1]["codes"] == codes:
                items[-1]["imgs"].append(img)
                continue
            items.append({
                "id": slug(cat["id"] + "-" + "-".join(codes)),
                "cat": cat["id"],
                "codes": codes,
                "note": note,
                "imgs": [img],
            })
        cats.append({k: cat[k] for k in ("id", "name", "tagline")} | {"items": items, "cover": items[0]["imgs"][0]["t"]})
    return {"categories": cats}


def ver(path):
    """Short content hash for links (?v=...): files can be cached for a year and still update when they change."""
    return hashlib.md5(Path(path).read_bytes()).hexdigest()[:8]


SMALL = {"s": 240, "m": 360}  # smaller copies of each thumbnail: s = rate list / search / home cards, m = category grid on phones


def version_images(catalogue):
    for c in catalogue["categories"]:
        for it in c["items"]:
            for im in it["imgs"]:
                thumb = OUT / im["t"].split("?")[0]
                for k, w in SMALL.items():
                    dst = thumb.with_name(thumb.stem[:-2] + f"-{k}.webp")
                    if not dst.exists() or dst.stat().st_mtime < thumb.stat().st_mtime:
                        t = Image.open(thumb)
                        t.resize((w, round(w * t.height / t.width)), Image.LANCZOS).save(dst, "WEBP", quality=70, method=6)
                    im[k] = dst.relative_to(OUT).as_posix()
                for k in ("t", "f", *SMALL):
                    path = im[k].split("?")[0]
                    im[k] = f"{path}?v={ver(OUT / path)}"
        c["cover"] = c["items"][0]["imgs"][0]["t"]
        c["coverS"] = c["items"][0]["imgs"][0]["s"]


def render_pdf_photos(catalogue):
    """img/<cat>/<page>-p.jpg (200 px JPEG) + data/photos.json: item code -> photo, for the bot's PDF order confirmation
    (PDF files can show JPEG, not WebP). Rate-sheet codes are linked through the sheet's \"Catalogue item\" column."""
    def keys(code):
        n = norm(code)
        c = re.sub(r"^(RC|RJ|RAJ|RT|TESLA|T)(?=\d)", "", n)
        z = lambda s: re.sub(r"(^|[A-Z])0+(?=\d)", r"\1", s)
        return [k for k in dict.fromkeys([n, c, z(n), z(c)]) if k]
    photos = {}
    for c in catalogue["categories"]:
        for it in c["items"]:
            thumb = OUT / it["imgs"][0]["t"].split("?")[0]
            dst = thumb.with_name(thumb.stem[:-2] + "-p.jpg")
            if not dst.exists() or dst.stat().st_mtime < thumb.stat().st_mtime:
                t = Image.open(thumb).convert("RGB")
                t.resize((200, round(200 * t.height / t.width)), Image.LANCZOS).save(dst, "JPEG", quality=72, optimize=True)
            path = dst.relative_to(OUT).as_posix() + f"?v={ver(dst)}"
            for code in it["codes"]:
                for k in keys(code):
                    photos.setdefault(k, path)
    for r in build_rates():
        link = norm(r.get("Catalogue item", ""))
        if link and link in photos:
            for k in keys(r.get("Code", "")):
                photos.setdefault(k, photos[link])
    logo = render_logo_pdf()
    return {"logo": f"assets/logo-pdf.jpg?v={ver(logo)}", "photos": photos}


def render_nav_icons(catalogue):
    """Small round photos for the category buttons in the header (52 px, shown at 26 px)."""
    icons = {}
    for c in catalogue["categories"]:
        im = Image.open(OUT / c["items"][0]["imgs"][0]["t"].split("?")[0]).convert("RGB")
        s = im.width
        top = round((im.height - s) * 0.6)
        dst = OUT / "assets" / f"nav-{c['id']}.webp"
        im.crop((0, top, s, top + s)).resize((52, 52), Image.LANCZOS).save(dst, "WEBP", quality=80, method=6)
        icons[c["id"]] = f"assets/nav-{c['id']}.webp?v={ver(dst)}"
    return icons


def minify(src, dst):
    """Minified copy made with bun (https://bun.sh) when it is installed, a plain copy otherwise."""
    try:
        subprocess.run(["bun", "build", str(src), "--minify", "--outfile", str(dst)], check=True, capture_output=True)
    except (OSError, subprocess.CalledProcessError):
        shutil.copy(src, dst)


def i18n_en():
    """The English texts from i18n.js (read with node), used to fill pages in before the script runs."""
    js = ("global.window={};require(process.argv[1]);const I=window.RC_I18N,t={},c={};"
          "for(const[k,v]of Object.entries(I.text))t[k]=v.en;for(const[k,v]of Object.entries(I.categories))c[k]=v.en;"
          "console.log(JSON.stringify({t,c}))")
    try:
        r = subprocess.run(["node", "-e", js, str(SRC / "assets" / "i18n.js")], capture_output=True, check=True)
        return json.loads(r.stdout.decode("utf8"))
    except (OSError, subprocess.CalledProcessError, ValueError):
        print("  node not found: pages will be filled in by the script")
        return None


def prerender(catalogue, en, nav):
    """HTML for the parts app.js would otherwise add after loading, so nothing jumps around (CLS)."""
    e = lambda v: html.escape(str(v), quote=True)
    out = {f"nav_{k.replace('-', '')}": v for k, v in nav.items()}
    out.update(promo="", catcards="")
    if not en:
        return out
    t, cats = en["t"], en["c"]
    def tr(key, **v):
        s = t.get(key, key)
        for k, x in v.items():
            s = s.replace("{" + k + "}", str(x))
        return s
    open_mov = next(x["mov"] for x in config.BUSINESS_TYPES if x.get("open") and not x.get("carton"))
    out["promo"] = f"<b>{e(tr('promo.title'))}</b> " + e(tr("promo.text", x=f"₹{open_mov:,}", b=config.CARTON["box"]["min"], p=config.CARTON["packet"]["min"]))
    FIRST, LAZY = 'fetchpriority="high"', 'loading="lazy"'
    out["catcards"] = "".join(
        f'<a class="cat-card" href="{{{{base}}}}{c["id"]}/"><img src="{{{{base}}}}{c["cover"]}" '
        f'srcset="{{{{base}}}}{c["coverS"]} 240w, {{{{base}}}}{c["cover"]} 480w" sizes="(min-width: 720px) 400px, 112px" alt="{e(cats[c["id"]][0])}" '
        f'{FIRST if i == 0 else LAZY} width="480" height="620">'
        f'<div><h3>{e(cats[c["id"]][0])}</h3><p>{e(cats[c["id"]][1])}</p><span class="count">{e(tr("items", n=len(c["items"])))}</span></div></a>'
        for i, c in enumerate(catalogue["categories"]))
    for c in catalogue["categories"]:
        name = cats[c["id"]][0]
        cards = []
        for i, it in enumerate(c["items"]):
            img, title = it["imgs"][0], " / ".join(it["codes"])
            load = FIRST if i == 0 else 'loading="eager"' if i < 4 else LAZY
            cards.append(
                f'<article class="card" id="{it["id"]}"><button class="card-img" type="button" data-open="{it["id"]}" aria-label="{e(title)}">'
                f'<img src="{{{{base}}}}{img["t"]}" srcset="{{{{base}}}}{img["m"]} 360w, {{{{base}}}}{img["t"]} 480w" sizes="(min-width: 980px) 260px, (min-width: 640px) 31vw, 48vw" '
                f'alt="{e(title)} – {e(name)}" width="480" height="{round(480 * img["h"] / img["w"])}" {load} decoding="async">'
                + (f'<span class="more">{e(tr("photos", n=len(it["imgs"])))}</span>' if len(it["imgs"]) > 1 else "")
                + f'</button><div class="card-body"><div class="card-code">{e(title)}</div>'
                + (f'<div class="card-note">{e(it["note"])}</div>' if it["note"] else "") + "</div></article>")
        out["grid_" + c["id"]] = "".join(cards)
    return out


def build_rates():
    """Fallback copy of the rate sheet (the live site reads the Google Sheet when configured)."""
    ws = openpyxl.load_workbook(RATES_XLSX, data_only=True)["Rates"]
    head = [str(c.value).strip() for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(head, r))
        d.pop("Note", None)
        rows.append({k: ("" if v is None else v) for k, v in d.items()})
    return rows


def render_pages(catalogue, fill):
    shared = (SRC / "partials" / "head.html").read_text(encoding="utf8")
    header = (SRC / "partials" / "header.html").read_text(encoding="utf8")
    footer = (SRC / "partials" / "footer.html").read_text(encoding="utf8")
    cfg = {
        "whatsapp": config.BOT_WHATSAPP,
        "phoneDisplay": config.PHONE_DISPLAY,
        "types": config.BUSINESS_TYPES,
        "carton": config.CARTON,
        "gst": config.GST_PCT,
    }

    def page(template, out_path, **vars):
        html = (SRC / template).read_text(encoding="utf8")
        depth = len(out_path.relative_to(OUT).parts) - 1
        base = "../" * depth if depth else "./"
        vars = {"head": shared, "header": header, "footer": footer, **fill, **vars}
        for _ in range(2):  # partials may contain placeholders too
            for k, v in vars.items():
                html = html.replace("{{" + k + "}}", str(v))
        html = html.replace("{{base}}", base).replace("{{config}}", json.dumps(cfg, ensure_ascii=False))
        html = html.replace("{{site}}", config.SITE_URL).replace("{{phone}}", config.PHONE_DISPLAY)
        left = re.findall(r"\{\{?[a-z_]+\}?\}", html.split("<style>")[0] + html.split("</style>")[-1])
        if left:
            raise SystemExit(f"{out_path}: placeholders not filled in: {sorted(set(left))}")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(html, encoding="utf8")

    page("index.html", OUT / "index.html",
         title="Raj Creation – Wholesale Hair Accessories",
         description="Manufacturer & wholesaler of claw clips, scrunchies, rubber bands, bun donuts and hair tie bows. Browse the catalogue and search any item code.",
         ogimage=catalogue["categories"][1]["items"][2]["imgs"][0]["f"], page="home", cat="", robots="index,follow")
    for c in catalogue["categories"]:
        page("category.html", OUT / c["id"] / "index.html",
             title=f"{c['name']} – Raj Creation Catalogue",
             description=f"{c['tagline']}. Wholesale catalogue with item codes – Raj Creation.",
             ogimage=c["items"][0]["imgs"][0]["f"], page="category", cat=c["id"], catname=c["name"], tagline=c["tagline"],
             grid=fill.get("grid_" + c["id"], ""),
             robots="index,follow")
    page("rates.html", OUT / "p" / config.PRICE_TOKEN / "index.html",
         title="Rate List – Raj Creation", description="Raj Creation wholesale rate list.",
         ogimage=catalogue["categories"][0]["items"][0]["imgs"][0]["f"], page="rates", cat="", robots="noindex,nofollow")


def main():
    force = "--images" in sys.argv
    OUT.mkdir(exist_ok=True)
    print("images…")
    sizes = render_images(force)
    render_logo()
    render_logo_mark()
    catalogue = build_catalogue(sizes)
    version_images(catalogue)
    nav = render_nav_icons(catalogue)
    (OUT / "data").mkdir(exist_ok=True)
    (OUT / "data" / "catalogue.json").write_text(json.dumps(catalogue, ensure_ascii=False, separators=(",", ":")), encoding="utf8")
    (OUT / "data" / "photos.json").write_text(json.dumps(render_pdf_photos(catalogue), separators=(",", ":")), encoding="utf8")

    price_dir = OUT / "p" / config.PRICE_TOKEN
    price_dir.mkdir(parents=True, exist_ok=True)
    (price_dir / "rates.json").write_text(json.dumps(build_rates(), ensure_ascii=False, separators=(",", ":")), encoding="utf8")
    (price_dir / "source.json").write_text(json.dumps({"csv": config.RATES_CSV_URL}), encoding="utf8")
    for old in (OUT / "p").iterdir():  # remove price pages for old tokens
        if old.is_dir() and old.name != config.PRICE_TOKEN:
            shutil.rmtree(old)

    (OUT / "assets" / "style.css").unlink(missing_ok=True)  # now inside each page
    minify(SRC / "assets" / "style.css", HERE / "_style.min.css")
    css = (HERE / "_style.min.css").read_text(encoding="utf8").strip()
    (HERE / "_style.min.css").unlink()
    for name in ("app.js", "i18n.js"):
        minify(SRC / "assets" / name, OUT / "assets" / name)
    shutil.copytree(SRC / "assets" / "fonts", OUT / "assets" / "fonts", dirs_exist_ok=True)
    fill = prerender(catalogue, i18n_en(), nav)
    fill["css"] = css
    for key, path in (("v_app", "assets/app.js"), ("v_i18n", "assets/i18n.js"), ("v_logo", "assets/logo-mark.webp"),
                      ("v_inter", "assets/fonts/inter-latin.woff2"), ("v_rupee", "assets/fonts/inter-rupee.woff2"),
                      ("v_playfair", "assets/fonts/playfair-latin.woff2")):
        fill[key] = ver(OUT / path)
    for name in (".htaccess", "robots.txt", "404.html"):
        if (SRC / name).exists():
            shutil.copy(SRC / name, OUT / name)
    (OUT / "chat").mkdir(exist_ok=True)  # /chat/?n=<number> forwards to WhatsApp (button in Mahek's alert template)
    shutil.copy(SRC / "chat" / "index.html", OUT / "chat" / "index.html")
    render_pages(catalogue, fill)

    n_items = sum(len(c["items"]) for c in catalogue["categories"])
    size_mb = sum(f.stat().st_size for f in OUT.rglob("*") if f.is_file()) / 1e6
    print(f"done: {n_items} catalogue items, site size {size_mb:.1f} MB -> {OUT}")
    print(f"price page: {config.SITE_URL}/p/{config.PRICE_TOKEN}/")


if __name__ == "__main__":
    main()
